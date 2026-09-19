#!/usr/bin/env python3
"""
Збирає сайт-словник із Excel-файлу.

    dictionary.xlsx  ->  site/data.js

Запуск:  python build.py
На GitHub це відбувається автоматично після кожного завантаження dictionary.xlsx.
"""
import datetime
import json
import pathlib
import re
import sys

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent
XLSX = ROOT / "dictionary.xlsx"
OUT = ROOT / "site" / "data.js"
SHEET = "Словник"
SRC_SHEET = "Джерела"

# Назви колонок в Excel -> внутрішні поля
HEADERS = {
    "лемма": "lemma",
    "граматика": "gram",
    "тип": "type",
    "англійська": "en",
    "українська": "uk",
    "примітка": "note",
    "приклади": "ex",
    "синоніми": "syn",
    "антоніми": "ant",
    "джерело": "src",
}

# Кириличні літери, візуально ідентичні латинським (типова помилка при наборі)
CYR2LAT = {
    "і": "i", "І": "I", "е": "e", "Е": "E", "о": "o", "О": "O", "с": "c", "С": "C",
    "а": "a", "А": "A", "р": "p", "Р": "P", "х": "x", "Х": "X", "у": "y", "У": "Y",
    "ў": "y", "Ў": "Y", "ё": "ë", "Ё": "Ë", "ї": "ï",
}
# Латинські літери, візуально ідентичні кириличним (у українських словах)
LAT2CYR = {
    "c": "с", "o": "о", "a": "а", "e": "е", "p": "р", "i": "і", "x": "х", "y": "у",
    "C": "С", "O": "О", "A": "А", "E": "Е", "P": "Р", "I": "І", "X": "Х",
    "H": "Н", "K": "К", "M": "М", "T": "Т", "B": "В",
}
CYR_RE = re.compile("[\u0400-\u04FF]")
LAT_RE = re.compile("[A-Za-z]")


# ---------------------------------------------------------------- очищення

def clean_ws(s):
    """Пробіли, нерозривні пробіли, м'які переноси, нульова ширина."""
    if s is None:
        return ""
    s = str(s)
    s = s.replace("\u00ad", "").replace("\u200b", "").replace("\ufeff", "")
    s = s.replace("\xa0", " ").replace("\u2009", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r" *\n *", "\n", s)
    return s.strip()


def fix_latin(s):
    """Кириличні двійники -> латинські (поза дужками; у дужках лише якщо
    там немає справжньої кирилиці, напр. '(гр.)' лишається як є)."""
    parts = re.split(r"(\([^)]*\))", s)
    out = []
    for p in parts:
        if p.startswith("("):
            cyr = CYR_RE.findall(p)
            if cyr and all(c in CYR2LAT for c in cyr):
                p = "".join(CYR2LAT.get(c, c) for c in p)
        else:
            p = "".join(CYR2LAT.get(c, c) for c in p)
        out.append(p)
    return "".join(out)


def fix_ukr(s):
    """Латинські двійники всередині українських слів -> кириличні;
    римські цифри в дужках -> латиницею."""
    def word(m):
        w = m.group(0)
        if CYR_RE.search(w) and LAT_RE.search(w):
            return "".join(LAT2CYR.get(c, c) for c in w)
        return w
    s = re.sub(r"[^\W\d_]+", word, s)
    s = re.sub(r"\(([ІIVXХ]+)\)", lambda m: "(" + m.group(1).replace("І", "I").replace("Х", "X") + ")", s)
    s = re.sub(r"[ʼ`´']", "’", s)
    return s


def split_top(s, seps):
    """Ділить рядок за роздільниками лише поза дужками."""
    out, depth, cur = [], 0, []
    for ch in s:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if depth == 0 and ch in seps:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return out


def has_top(s, ch):
    return len(split_top(s, ch)) > 1


def split_latin(s):
    """'abdomen, inis, n.' -> ('abdomen', 'inis, n.')"""
    if has_top(s, ";"):
        return s, ""
    parts = split_top(s, ",")
    if all(p.strip().endswith("-") for p in parts):  # префікси: ante-, prae-
        return s, ""
    depth, cut = 0, None
    for i, ch in enumerate(s):
        if ch == "(":
            if depth == 0 and i > 0 and s[i - 1] == " ":
                cut = i
                break
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        elif ch == "," and depth == 0:
            cut = i
            break
    if cut is None:
        return s, ""
    lemma = s[:cut].strip()
    gram = s[cut:].lstrip(", ").strip()
    if " " in lemma and s[cut] == "," and not re.search(r"\b[mfn]\.", gram) and not gram.startswith("("):
        return "; ".join(x.strip() for x in split_top(s, ",")), ""  # кілька рівнозначних назв
    m = re.fullmatch(r"(.+?) ([mfn]\.)", lemma)
    if m:
        lemma, gram = m.group(1), (m.group(2) + " " + gram).strip()
    return lemma, gram


def classify(lemma, gram):
    parts = re.split(r"[;,]\s*", lemma)
    if all(p.endswith("-") for p in parts):
        return "префікс"
    if not gram and re.fullmatch(r"[A-Z][a-zë]+ [a-z][a-z-]+", lemma):
        return "ботанічна назва"
    if re.search(r"\b[mfn]\.", gram):
        return "іменник"
    if re.fullmatch(r"\([IVX]+\)", gram) and " " not in lemma:
        return "числівник"
    if re.search(r"\b\w+(avi|ui|ivi|si|xi|di|ti|vi), \w+um, (are|ere|ire)$", gram):
        return "дієслово"
    if gram and " " not in lemma:
        return "прикметник"
    if " " in lemma or ";" in lemma:
        return "словосполучення"
    return "інше"


def parse_links(cell):
    """'pulsus parvus; os, ossis, n.' -> ['pulsus parvus', 'os, ossis, n.']"""
    return [fix_latin(x.strip()) for x in split_top(clean_ws(cell), ";") if x.strip()]


def norm_key(s):
    return re.sub(r"\s+", " ", s.lower()).strip()


def resolve_links(entries, warnings):
    """Зіставляє посилання з гаслами: за повною лемою або лема + граматика.
    Зв'язки НЕ дзеркалюються автоматично; висячі, неоднозначні й односторонні — у звіті."""
    by_lemma, by_full = {}, {}
    for e in entries:
        by_lemma.setdefault(norm_key(e["l"]), []).append(e)
        by_full[norm_key(f"{e['l']}, {e['g']}" if e["g"] else e["l"])] = e

    def lookup(e, raw, label):
        k = norm_key(raw)
        target = by_full.get(k)
        if target is None:
            cands = by_lemma.get(k, [])
            if len(cands) == 1:
                target = cands[0]
            elif len(cands) > 1:
                hint = "; ".join(f"«{c['l']}, {c['g']}»" for c in cands)
                warnings.append(f"«{e['l']}»: {label} «{raw}» неоднозначний — вкажіть граматику: {hint}")
                return None
        if target is None:
            warnings.append(f"«{e['l']}»: {label} «{raw}» — такого гасла немає (висяче посилання)")
        elif target is e:
            warnings.append(f"«{e['l']}»: {label} посилається сам на себе")
            return None
        return target

    for src, dst, label in (("_sy", "sy", "синонім"), ("_an", "an", "антонім")):
        targets = {}
        for e in entries:
            out, tg = [], []
            for raw in e[src]:
                t = lookup(e, raw, label)
                if t is not None:
                    out.append({"l": t["l"], "g": t["g"], "s": t["s"]})
                    tg.append(t)
                elif warnings and "висяче посилання" in warnings[-1] and f"«{raw}»" in warnings[-1]:
                    out.append({"l": raw, "g": "", "s": ""})  # показуємо текстом без посилання
            e[dst] = out
            targets[id(e)] = tg
        for e in entries:
            for t in targets[id(e)]:
                if e not in targets[id(t)]:
                    warnings.append(f"односторонній {label}: «{e['l']}» → «{t['l']}» (зворотного зв'язку немає)")
    for e in entries:
        e.pop("_sy", None); e.pop("_an", None)


# ---------------------------------------------------------------- збірка

def read_rows():
    if not XLSX.exists():
        sys.exit("Не знайдено dictionary.xlsx поруч із build.py")
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"У файлі немає аркуша «{SHEET}»")
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    head = [clean_ws(c).lower() for c in next(rows)]
    idx = {HEADERS[h]: i for i, h in enumerate(head) if h in HEADERS}
    missing = [h for h in ("лемма", "українська") if HEADERS[h] not in idx]
    if missing:
        sys.exit("У шапці аркуша «Словник» немає обов'язкових колонок: " + ", ".join(missing))
    for r in rows:
        rec = {k: (r[i] if i < len(r) else None) for k, i in idx.items()}
        if not clean_ws(rec.get("lemma")):
            continue
        yield rec


def read_sources():
    """Аркуш «Джерела»: № | Джерело | Внесок | Посилання. Необов'язковий."""
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if SRC_SHEET not in wb.sheetnames:
        return []
    rows = wb[SRC_SHEET].iter_rows(values_only=True)
    head = [clean_ws(c).lower() for c in next(rows)]
    def col(name):
        return head.index(name) if name in head else None
    ci = {k: col(k) for k in ("№", "джерело", "внесок", "посилання")}
    out = []
    for r in rows:
        get = lambda k: clean_ws(r[ci[k]]) if ci[k] is not None and ci[k] < len(r) else ""
        code = get("№")
        if code.endswith(".0"):
            code = code[:-2]
        if not code or not get("джерело"):
            continue
        out.append({"n": code, "t": get("джерело"), "c": get("внесок"), "u": get("посилання")})
    return out


def parse_examples(cell):
    ex = []
    for line in clean_ws(cell).split("\n"):
        line = line.strip()
        if not line:
            continue
        if "//" in line:
            t, s = line.split("//", 1)
            ex.append({"t": t.strip(), "s": s.strip()})
        else:
            ex.append({"t": line, "s": ""})
    return ex


def build():
    entries, warnings, seen = [], [], {}
    for n, rec in enumerate(read_rows(), start=2):
        lemma = fix_latin(clean_ws(rec["lemma"]))
        gram = fix_latin(clean_ws(rec.get("gram")))
        if not gram and "," in lemma and not has_top(lemma, ";"):
            lemma, gram = split_latin(lemma)  # лемму вписали разом із граматикою
        uk = [x.strip() for x in split_top(fix_ukr(clean_ws(rec.get("uk"))), ";") if x.strip()]
        en = [x.strip() for x in split_top(clean_ws(rec.get("en")), ";") if x.strip()]
        typ = clean_ws(rec.get("type")) or classify(lemma, gram)
        ent = {
            "l": lemma, "g": gram, "t": typ, "u": uk, "e": en,
            "n": clean_ws(rec.get("note")), "x": parse_examples(rec.get("ex")),
            "_sy": parse_links(rec.get("syn")), "_an": parse_links(rec.get("ant")),
        }
        codes = []
        for c in split_top(clean_ws(rec.get("src")), ";"):
            c = c.strip()
            c = c[:-2] if c.endswith(".0") else c
            if c and c not in codes:
                codes.append(c)
        if codes:
            ent["r"] = codes
        if not uk and not en:
            warnings.append(f"рядок {n}: «{lemma}» — немає жодного перекладу")
        key = (lemma, gram)
        if key in seen:  # той самий запис двічі — зливаємо переклади
            old = seen[key]
            old["u"] += [x for x in uk if x not in old["u"]]
            old["e"] += [x for x in en if x not in old["e"]]
            old["r"] = old.get("r", []) + [x for x in ent.get("r", []) if x not in old.get("r", [])]
            old["_sy"] += [x for x in ent["_sy"] if x not in old["_sy"]]
            old["_an"] += [x for x in ent["_an"] if x not in old["_an"]]
            warnings.append(f"рядок {n}: «{lemma}» уже є — переклади об'єднано")
            continue
        seen[key] = ent
        entries.append(ent)

    # стабільні адреси сторінок: із самої леми
    used = {}
    for e in entries:
        base = re.sub(r"[^a-z0-9]+", "-", e["l"].lower().split(";")[0]).strip("-") or "x"
        used[base] = used.get(base, 0) + 1
        e["s"] = base if used[base] == 1 else f"{base}-{used[base]}"

    resolve_links(entries, warnings)

    sources = read_sources()
    known = {x["n"] for x in sources}
    if len({x["n"] for x in sources}) != len(sources):
        warnings.append("аркуш «Джерела»: номери джерел повторюються")
    for e in entries:
        for c in e.get("r", []):
            if c not in known:
                warnings.append(f"«{e['l']}»: джерело «{c}» не описане на аркуші «Джерела»")
    for x in sources:
        x["k"] = sum(1 for e in entries if x["n"] in e.get("r", []))

    data = {
        "built": datetime.date.today().isoformat(),
        "entries": entries,
        "sources": sources,
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(
        "window.DICT=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    en_count = sum(1 for e in entries if e["e"])
    print(f"Готово: {len(entries)} гасел, з англійським перекладом — {en_count}, джерел — {len(sources)}.")
    for w in warnings:
        print("Увага:", w)


if __name__ == "__main__":
    build()
